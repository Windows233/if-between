from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent
SOURCE = ROOT / "2026-10-25_D4_教室同步与路线阻断策划.md"
OUTPUT = ROOT / "2026-10-25_D4_教室同步与路线阻断策划.docx"


def set_cell(cell, text, header=False):
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(9 if not header else 9.5)
            run.font.bold = header
            if header:
                run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(0.72), Inches(3.0), Inches(3.0)] if len(headers) == 3 else None
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, True)
        shading = table.rows[0].cells[i]._tc.get_or_add_tcPr()
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        fill = OxmlElement("w:shd")
        fill.set(qn("w:fill"), "245B68")
        shading.append(fill)
        if widths:
            table.columns[i].width = widths[i]
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if widths:
                cells[i].width = widths[i]
            if row_index % 2:
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn
                shading = cells[i]._tc.get_or_add_tcPr()
                fill = OxmlElement("w:shd")
                fill.set(qn("w:fill"), "F1F5F4")
                shading.append(fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table(lines):
    values = []
    for line in lines:
        if not line.strip() or line.strip().startswith("| ---"):
            continue
        values.append([cell.strip().replace("`", "") for cell in line.strip().strip("|").split("|")])
    return values


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(39, 49, 52)
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size, color in (("Heading 1", 16, "245B68"), ("Heading 2", 12, "2D6A4F")):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(5)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(line[2:])
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(36, 91, 104)
            index += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
            index += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
            index += 1
            continue
        if line.startswith("|"):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            rows = parse_table(block)
            if rows:
                add_table(doc, rows[0], rows[1:])
            continue
        paragraph = doc.add_paragraph()
        text = line.replace("**", "").replace("`", "")
        paragraph.add_run(text)
        index += 1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("IF：如果之间  |  D4团队执行稿")
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(105, 119, 120)
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
